from pathlib import Path
import re
import json
from collections import Counter
from openpyxl import load_workbook

ROOT = Path('/home/ubuntu/Careerbridge-data/Careerbridge(data)')
APP = Path('/home/ubuntu/MCA-project-careerBridge/src/app')
DATA = Path('/home/ubuntu/MCA-project-careerBridge/src/data')

DATASETS = {
    'institutions': ('Nagaland.xlsx', 'Master Census', ['institution_code', 'name', 'type', 'ownership', 'district', 'verification_status', 'source_url', 'last_verified_at']),
    'exams': ('EntranceExam.xlsx', 'All Exams', ['exam_slug', 'name', 'short_name', 'category', 'scope', 'state', 'conducting_body', 'applies_to', 'eligibility', 'level_stage', 'exam_mode', 'frequency', 'preparation', 'official_website', 'status', 'verification_status', 'source_url', 'last_verified_at']),
    'scholarships': ('scholarships.xlsx', 'Master', ['scholarship_slug', 'name', 'provider', 'category', 'eligibility', 'amount_note', 'deadline_note', 'documents', 'official_url', 'applies_to_stage', 'verification_status', 'source_url', 'last_verified_at']),
}


def norm(v):
    return '' if v is None else str(v).strip()


def validate_workbook(kind, spec):
    filename, sheet_name, required = spec
    path = ROOT / filename
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header_index = next(i for i, row in enumerate(rows) if any(norm(x).lower() in {'institution code', 'exam_slug', 'scholarship_slug'} for x in row))
    raw_headers = [norm(x) for x in rows[header_index]]
    header_aliases = {'Institution Code': 'institution_code', 'Name': 'name', 'Type': 'type', 'Ownership': 'ownership', 'District': 'district', 'Verification Status': 'verification_status', 'Source Url': 'source_url', 'Last Verified At': 'last_verified_at'}
    headers = [header_aliases.get(h, h) for h in raw_headers]
    positions = {h: i for i, h in enumerate(headers)}
    missing = [h for h in required if h not in positions]
    records = [dict(zip(headers, row)) for row in rows[header_index + 1:] if any(norm(x) for x in row)]
    key = required[0]
    keys = [norm(r.get(key)) for r in records]
    duplicates = sorted(k for k, n in Counter(keys).items() if k and n > 1)
    blank_counts = {h: sum(not norm(r.get(h)) for r in records) for h in required}
    urls = []
    for r in records:
        for field in ('official_website', 'official_url', 'source_url'):
            value = norm(r.get(field))
            if value:
                urls.append((field, value))
    bad_urls = [(f, u) for f, u in urls if f in ('official_website', 'official_url') and not re.match(r'^https?://', u)]
    source_annotations = [(f, u) for f, u in urls if f == 'source_url' and not re.match(r'^https?://', u)]
    dates = [norm(r.get('last_verified_at')) for r in records]
    return {
        'file': filename,
        'sheet': sheet_name,
        'rows': len(records),
        'columns': len(headers),
        'missing_headers': missing,
        'duplicate_keys': duplicates,
        'blank_required_fields': {k: v for k, v in blank_counts.items() if v},
        'bad_urls': bad_urls[:10],
        'source_annotations_normalized_by_importer': source_annotations[:10],
        'verification_dates': sorted(set(dates)),
        'districts': sorted({norm(r.get('district')) for r in records if norm(r.get('district'))}),
        'categories': sorted({norm(r.get('category')) for r in records if norm(r.get('category'))}),
        'stages': sorted({stage.strip() for r in records for stage in norm(r.get('applies_to_stage')).split(';') if stage.strip()}),
    }


def source_checks():
    institution_page = (APP / 'institutions/page.tsx').read_text()
    institution_detail = (APP / 'institutions/[code]/page.tsx').read_text()
    exam_page = (APP / 'exams/page.tsx').read_text()
    scholarship_page = (APP / 'scholarships/page.tsx').read_text()
    institution_data = (DATA / 'institutions.ts').read_text()
    return {
        'institutions': {
            'required_ui_tokens': all(token in institution_page + institution_detail for token in ['institution.name', 'institution.city', 'institution.district', 'institution.officialWebsite', 'institution.verificationStatus']),
            'district_count_in_ui_seed': len(re.findall(r'level: "district"', institution_data)),
            'search_and_filters_present': all(token in institution_page for token in ['Search institutions', 'More filters', 'district', 'ownership', 'level']),
        },
        'exams': {
            'required_ui_tokens': all(token in exam_page for token in ['exam.name', 'exam.conductingBody', 'exam.eligibility', 'exam.preparation', 'exam.officialWebsite', 'exam.sourceUrl', 'exam.verificationStatus']),
            'metadata_displayed': all(token in exam_page for token in ['exam.category', 'exam.scope', 'exam.state', 'exam.examMode', 'exam.status']),
        },
        'scholarships': {
            'required_ui_tokens': all(token in scholarship_page for token in ['scholarship.name', 'scholarship.provider', 'scholarship.eligibility', 'scholarship.amountNote', 'scholarship.deadlineNote', 'scholarship.officialUrl', 'scholarship.sourceUrl', 'scholarship.verificationStatus']),
            'all_stage_filters_present': all(stage in scholarship_page for stage in ['class9', 'class10', 'class11', 'class12', 'diploma', 'undergraduate', 'postgraduate']),
        },
    }

result = {'workbooks': {kind: validate_workbook(kind, spec) for kind, spec in DATASETS.items()}, 'ui_source_checks': source_checks()}
print(json.dumps(result, ensure_ascii=False, indent=2))

failures = []
for kind, report in result['workbooks'].items():
    for field in ('missing_headers', 'duplicate_keys', 'blank_required_fields', 'bad_urls'):
        if report[field]: failures.append(f'{kind}.{field}')
for kind, checks in result['ui_source_checks'].items():
    for name, passed in checks.items():
        if not passed: failures.append(f'ui.{kind}.{name}')
print(f'VALIDATION_STATUS={"PASS" if not failures else "FAIL"}')
print('FAILURES=' + (','.join(failures) if failures else 'none'))
